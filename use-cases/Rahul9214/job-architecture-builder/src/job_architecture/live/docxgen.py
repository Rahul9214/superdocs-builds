"""Deterministic DOCX generation. This is the only module that imports python-docx."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

from job_architecture.documents import render_framework_markdown, render_profile_markdown
from job_architecture.fixtures import load_corpus
from job_architecture.framework import FrameworkDocument
from job_architecture.live.subset import DemoSubset, load_demo_subset, subset_profiles
from job_architecture.models import RoleProfile


def markdown_to_docx(markdown: str, destination: Path, *, title: str | None = None) -> Path:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    first = True
    for raw_line in markdown.replace("\r\n", "\n").split("\n"):
        line = raw_line.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=0 if first else 1)
            first = False
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
            first = False
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
            continue
        if line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
            continue
        document.add_paragraph(line)
    if title:
        document.core_properties.title = title
        document.core_properties.author = "Job Architecture Builder"
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))
    return destination


def framework_template_markdown() -> str:
    return "\n".join(
        [
            "# Job architecture framework",
            "",
            "## Purpose",
            "[Insert approved purpose. Do not invent families or levels.]",
            "",
            "## Principles",
            "[Insert approved principles.]",
            "",
            "## Career tracks",
            "[Insert Individual Contributor and People Manager tracks.]",
            "",
            "## Canonical levels",
            "[Insert IC1–IC5 and M1–M3 with versioned dimensions.]",
            "",
            "## Job families",
            "[Insert only families from the approved architecture.]",
            "",
            "## Competency matrices",
            "[Insert constrained family matrices.]",
            "",
            "## Role mappings",
            "[Insert the demo-subset mappings only.]",
            "",
            "## Provisional roles",
            "[Insert provisional review notes.]",
            "",
            "## Misfits",
            "[None in the live demo subset unless listed.]",
            "",
        ]
    )


def role_profile_template_markdown() -> str:
    return "\n".join(
        [
            "# Role profile",
            "",
            "## Role title",
            "[title]",
            "",
            "## Job family",
            "[family]",
            "",
            "## Career track",
            "[track]",
            "",
            "## Level",
            "[level]",
            "",
            "## Classification",
            "[strong_fit or provisional]",
            "",
            "## Role purpose",
            "[purpose]",
            "",
            "## Responsibilities",
            "[responsibilities]",
            "",
            "## Scope / decision making",
            "[scope]",
            "",
            "## Core competencies",
            "[competencies]",
            "",
            "## Level expectations",
            "[canonical level-definition dimensions only]",
            "",
            "## Progression",
            "[next-level framing]",
            "",
            "## Source / evidence note",
            "[evidence note]",
            "",
        ]
    )


def write_demo_artifacts(
    destination: Path,
    *,
    subset: DemoSubset | None = None,
    framework: FrameworkDocument | None = None,
    profile: RoleProfile | None = None,
) -> dict[str, Path]:
    spec = subset or load_demo_subset()
    corpus = load_corpus(spec.corpus_id)
    destination.mkdir(parents=True, exist_ok=True)
    written: dict[str, Path] = {}
    for role in spec.roles:
        parsed = corpus.job_description(role.role_id)
        path = destination / f"{role.role_id}.docx"
        markdown_to_docx(parsed.raw_markdown, path, title=parsed.title)
        written[role.role_id] = path
    written["framework-template"] = markdown_to_docx(
        framework_template_markdown(),
        destination / "framework-template.docx",
        title="Job architecture framework template",
    )
    written["role-profile-template"] = markdown_to_docx(
        role_profile_template_markdown(),
        destination / "role-profile-template.docx",
        title="Role profile template",
    )
    if framework is not None:
        written["framework-starter"] = markdown_to_docx(
            framework_template_markdown(),
            destination / "framework-starter.docx",
            title="Job architecture framework",
        )
        written["framework-filled"] = markdown_to_docx(
            render_framework_markdown(framework),
            destination / "framework-filled.docx",
            title="Job architecture framework",
        )
        for item in subset_profiles(framework, spec):
            written[item.id] = markdown_to_docx(
                render_profile_markdown(item),
                destination / f"{item.id}.docx",
                title=item.display_title,
            )
    if profile is not None and profile.id not in written:
        written["profile-filled"] = markdown_to_docx(
            render_profile_markdown(profile),
            destination / f"{profile.id}.docx",
            title=profile.display_title,
        )
    return written
