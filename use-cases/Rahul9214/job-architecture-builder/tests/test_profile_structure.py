"""Structural profile validation and the duplicate level-expectations bug class."""

from __future__ import annotations

from pathlib import Path

from job_architecture.graph import DependencyGraphError
from job_architecture.live.docxgen import write_demo_artifacts
from job_architecture.live.instructions import (
    compact_profile_instruction,
    profile_duplicate_repair_instruction,
    subset_profile_for_demo,
)
from job_architecture.live.preservation import extract_docx_text
from job_architecture.live.subset import build_demo_framework, load_demo_subset
from job_architecture.profile_structure import (
    duplicate_block_via_narrow_header_replace,
    validate_level_expectations_block,
    validate_profile_document_text,
)
from job_architecture.propagation import analyze_level_change, plan_level_updates
from job_architecture.live.instructions import surgical_level_change


def _demo_profile():
    subset = load_demo_subset()
    _, framework = build_demo_framework(subset)
    return subset_profile_for_demo(framework, subset), framework, subset


def test_valid_generated_profile_has_each_level_expectation_field_once(tmp_path: Path) -> None:
    profile, _framework, subset = _demo_profile()
    block_report = validate_level_expectations_block(profile.level_expectations)
    assert block_report.ok, block_report.violations
    assert block_report.field_counts["header"] == 1
    assert block_report.field_counts["scope"] == 1
    assert block_report.field_counts["autonomy"] == 1
    assert block_report.field_counts["complexity"] == 1
    assert block_report.field_counts["impact"] == 1
    assert block_report.field_counts["leadership"] == 1
    assert block_report.field_counts["people_management"] == 1
    written = write_demo_artifacts(tmp_path, subset=subset, framework=_framework, profile=profile)
    text = extract_docx_text(written[profile.id])
    doc_report = validate_profile_document_text(text)
    assert doc_report.ok, doc_report.violations
    assert doc_report.section_counts["Level expectations"] == 1


def test_narrow_header_replace_with_full_block_duplicates_dimensions() -> None:
    profile, _framework, _subset = _demo_profile()
    original = validate_level_expectations_block(profile.level_expectations)
    assert original.ok
    duplicated = duplicate_block_via_narrow_header_replace(profile.level_expectations)
    report = validate_level_expectations_block(duplicated)
    assert not report.ok
    duplicate_fields = [item for item in report.violations if item.startswith("Duplicate")]
    assert duplicate_fields
    assert report.field_counts["scope"] == 2
    assert report.field_counts["complexity"] == 2
    assert report.field_counts["header"] == 1
    assert report.field_counts["autonomy"] == 2


def test_fill_instruction_against_filled_block_is_the_unsafe_span_class() -> None:
    profile, _framework, _subset = _demo_profile()
    instruction = compact_profile_instruction(profile)
    after = profile.level_expectations
    before = after.splitlines()[0]
    assert before in instruction
    result = duplicate_block_via_narrow_header_replace(after)
    assert validate_level_expectations_block(result).ok is False


def test_repair_instruction_contains_canonical_block_once() -> None:
    profile, _framework, _subset = _demo_profile()
    instruction = profile_duplicate_repair_instruction(profile.level_expectations)
    assert instruction.count(profile.level_expectations) == 1
    assert "duplicated" in instruction.lower()
    assert "exactly" in instruction.lower()
    assert validate_level_expectations_block(profile.level_expectations).ok


def test_plan_level_updates_rejects_duplicated_baseline() -> None:
    profile, framework, subset = _demo_profile()
    duplicated = duplicate_block_via_narrow_header_replace(profile.level_expectations)
    broken = profile.with_section("level_expectations", duplicated)
    old = framework.level(profile.level_id)
    new = surgical_level_change(old)
    profiles = tuple(broken if item.id == broken.id else item for item in framework.profiles)
    analysis = analyze_level_change(
        old,
        new,
        framework.dependency_graph,
        profiles,
        levels=framework.levels,
    )
    try:
        plan_level_updates(
            analysis,
            old_level=old,
            new_level=new,
            dependency_graph=framework.dependency_graph,
            profiles=profiles,
        )
    except DependencyGraphError as exc:
        assert "malformed" in str(exc).lower()
        return
    raise AssertionError("expected DependencyGraphError for duplicated baseline")


def test_missing_and_duplicate_section_headings_are_violations() -> None:
    missing = validate_profile_document_text("Role purpose\nHello\n")
    assert not missing.ok
    assert any("Level expectations" in item for item in missing.violations)
    doubled = "\n".join(
        [
            "Level expectations",
            "IC4 (definition version 1)",
            "Scope: one",
            "Autonomy / decision authority: one",
            "Complexity: one",
            "Impact: one",
            "Leadership: one",
            "People-management expectations: none",
            "Level expectations",
            "IC4 (definition version 1)",
            "Scope: one",
            "Autonomy / decision authority: one",
            "Complexity: one",
            "Impact: one",
            "Leadership: one",
            "People-management expectations: none",
        ]
    )
    report = validate_profile_document_text(doubled)
    assert not report.ok
    assert any("duplicate profile section" in item.lower() for item in report.violations)


_CANONICAL_FIELD_LINES = (
    "IC4 (definition version 1)",
    "Scope: Cross-team domain ownership",
    "Autonomy / decision authority: Decides independently within policy",
    "Complexity: High ambiguity across systems",
    "Impact: Org-wide product outcomes",
    "Leadership: Mentors other ICs",
    "People-management expectations: none",
)
_CANONICAL_VALUES = {
    "header": "IC4 (definition version 1)",
    "scope": "Cross-team domain ownership",
    "autonomy": "Decides independently within policy",
    "complexity": "High ambiguity across systems",
    "impact": "Org-wide product outcomes",
    "leadership": "Mentors other ICs",
    "people_management": "none",
}


def _wrapped_profile(body: str) -> str:
    return "\n".join(["Level expectations", body, "Progression", "Remain at IC4 or progress to IC5."])


def _assert_each_field_once(report) -> None:
    assert report.field_counts["header"] == 1
    assert report.field_counts["scope"] == 1
    assert report.field_counts["autonomy"] == 1
    assert report.field_counts["complexity"] == 1
    assert report.field_counts["impact"] == 1
    assert report.field_counts["leadership"] == 1
    assert report.field_counts["people_management"] == 1
    assert report.field_values == _CANONICAL_VALUES


def test_a_one_paragraph_per_field_is_valid() -> None:
    body = "\n".join(_CANONICAL_FIELD_LINES)
    report = validate_level_expectations_block(body)
    assert report.ok, report.violations
    _assert_each_field_once(report)
    doc = validate_profile_document_text(_wrapped_profile(body))
    assert doc.ok, doc.violations
    assert doc.section_counts["Level expectations"] == 1
    _assert_each_field_once(doc)


def test_b_concatenated_single_paragraph_is_valid() -> None:
    body = "".join(_CANONICAL_FIELD_LINES)
    assert "\n" not in body
    report = validate_level_expectations_block(body)
    assert report.ok, report.violations
    _assert_each_field_once(report)
    doc = validate_profile_document_text(_wrapped_profile(body))
    assert doc.ok, doc.violations
    _assert_each_field_once(doc)


def test_c_multiple_runs_and_breaks_in_one_paragraph(tmp_path: Path) -> None:
    from docx import Document

    from job_architecture.live.preservation import inspect_docx_layout

    path = tmp_path / "runs-one-paragraph.docx"
    document = Document()
    document.add_paragraph("Level expectations")
    paragraph = document.add_paragraph()
    for index, line in enumerate(_CANONICAL_FIELD_LINES):
        paragraph.add_run(line)
        if index < len(_CANONICAL_FIELD_LINES) - 1:
            paragraph.add_run().add_break()
    document.add_paragraph("Progression")
    document.add_paragraph("Remain at IC4 or progress to IC5.")
    document.save(str(path))
    layout = inspect_docx_layout(path)
    report = validate_profile_document_text(layout.text, layout=layout.as_dict())
    assert report.ok, report.violations
    _assert_each_field_once(report)
    assert layout.level_expectations_paragraphs == 1
    assert layout.level_expectations_runs == 13
    assert layout.level_expectations_breaks == 6
    assert report.layout["level_expectations_paragraphs"] == 1
    assert report.layout["level_expectations_runs"] == 13


def test_d_duplicate_scope_marker_is_rejected() -> None:
    lines = list(_CANONICAL_FIELD_LINES)
    lines.insert(2, "Scope: extra copy")
    report = validate_level_expectations_block("\n".join(lines))
    assert not report.ok
    assert report.field_counts["scope"] == 2
    assert any("Duplicate level-expectations field scope" in item for item in report.violations)
    concatenated = "".join(_CANONICAL_FIELD_LINES).replace("Scope:", "Scope: extra Scope:", 1)
    dup = validate_level_expectations_block(concatenated)
    assert not dup.ok
    assert dup.field_counts["scope"] == 2


def test_e_missing_complexity_marker_is_rejected() -> None:
    lines = [item for item in _CANONICAL_FIELD_LINES if not item.startswith("Complexity:")]
    report = validate_level_expectations_block("\n".join(lines))
    assert not report.ok
    assert report.field_counts["complexity"] == 0
    assert any("Missing required level-expectations field complexity" in item for item in report.violations)
    concatenated = "".join(item for item in _CANONICAL_FIELD_LINES if not item.startswith("Complexity:"))
    missing = validate_level_expectations_block(concatenated)
    assert not missing.ok
    assert missing.field_counts["complexity"] == 0


def test_f_duplicate_complete_block_is_rejected() -> None:
    block = "\n".join(_CANONICAL_FIELD_LINES)
    report = validate_level_expectations_block(f"{block}\n{block}")
    assert not report.ok
    assert report.field_counts["header"] == 2
    assert report.field_counts["scope"] == 2
    assert any("Duplicate complete Level expectations block" in item for item in report.violations)
    doubled_sections = "\n".join(
        [
            "Level expectations",
            block,
            "Level expectations",
            block,
        ]
    )
    doc = validate_profile_document_text(doubled_sections)
    assert not doc.ok
    assert doc.section_counts["Level expectations"] == 2


def test_g_field_values_remain_extractable_exactly() -> None:
    line_report = validate_level_expectations_block("\n".join(_CANONICAL_FIELD_LINES))
    concat_report = validate_level_expectations_block("".join(_CANONICAL_FIELD_LINES))
    assert line_report.field_values == _CANONICAL_VALUES
    assert concat_report.field_values == _CANONICAL_VALUES
    assert line_report.field_values == concat_report.field_values


def test_unknown_content_between_header_and_scope_is_rejected() -> None:
    body = "\n".join(
        [
            "IC4 (definition version 1)",
            "INSERTED JUNK LINE",
            *_CANONICAL_FIELD_LINES[1:],
        ]
    )
    report = validate_level_expectations_block(body)
    assert not report.ok
    assert any("Unexpected extra content between Level expectations fields" in item for item in report.violations)
